import os
import io
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import tempfile
import base64

import PyPDF2
from PIL import Image
from docx import Document
from docx.document import Document as DocumentType
import openpyxl

class ImageExtractor:
    """Service for extracting images from various document types"""
    
    def __init__(self):
        self.supported_formats = ['.pdf', '.docx', '.xlsx']
        self.image_counter = 0
    
    def extract_images_from_document(self, file_path: str, output_folder: str, 
                                   questions_with_images: List[Dict]) -> List[Dict]:
        """
        Extract images from document and save them with question IDs
        
        Args:
            file_path: Path to the source document
            output_folder: Folder to save extracted images
            questions_with_images: List of questions that have image references
        
        Returns:
            List of dictionaries with image info and file paths
        """
        file_extension = Path(file_path).suffix.lower()
        
        try:
            if file_extension == '.pdf':
                return self._extract_from_pdf(file_path, output_folder, questions_with_images)
            elif file_extension == '.docx':
                return self._extract_from_docx(file_path, output_folder, questions_with_images)
            elif file_extension == '.xlsx':
                return self._extract_from_xlsx(file_path, output_folder, questions_with_images)
            else:
                return []
        except Exception as e:
            print(f"Error extracting images: {str(e)}")
            return []
    
    def _extract_from_pdf(self, file_path: str, output_folder: str, 
                         questions_with_images: List[Dict]) -> List[Dict]:
        """Extract images from PDF - simplified version without PyMuPDF"""
        extracted_images = []
        
        try:
            # Note: Basic PDF image extraction is complex without PyMuPDF
            # For now, we'll create placeholder images if questions are detected
            print("PDF image extraction requires PyMuPDF. Creating placeholder images for detected questions.")
            
            for i, question in enumerate(questions_with_images):
                question_id = question.get("question_id", f"Q{i + 1:04d}")
                
                # Create a simple placeholder image
                placeholder_img = Image.new('RGB', (400, 300), color='lightgray')
                image_filename = f"{question_id}.png"
                image_path = os.path.join(output_folder, image_filename)
                
                placeholder_img.save(image_path)
                
                extracted_images.append({
                    "question_id": question_id,
                    "image_path": image_path,
                    "page_number": 1,
                    "image_index": i,
                    "format": "png",
                    "note": "Placeholder image - PDF image extraction requires PyMuPDF"
                })
                
        except Exception as e:
            print(f"Error creating placeholder images: {str(e)}")
        
        return extracted_images
    
    def _extract_from_docx(self, file_path: str, output_folder: str,
                          questions_with_images: List[Dict]) -> List[Dict]:
        """Extract images from Word document"""
        extracted_images = []
        
        try:
            doc = Document(file_path)
            image_counter = 0
            
            # Extract inline images
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    image_counter += 1
                    
                    # Get image data
                    image_data = rel.target_part.blob
                    
                    # Determine file extension
                    content_type = rel.target_part.content_type
                    if "png" in content_type:
                        ext = "png"
                    elif "jpeg" in content_type or "jpg" in content_type:
                        ext = "jpg"
                    else:
                        ext = "png"  # Default
                    
                    # Assign to question
                    question_id = self._assign_image_to_question(
                        0, image_counter, questions_with_images
                    )
                    
                    # Save image
                    image_filename = f"{question_id}.{ext}"
                    image_path = os.path.join(output_folder, image_filename)
                    
                    with open(image_path, "wb") as img_file:
                        img_file.write(image_data)
                    
                    extracted_images.append({
                        "question_id": question_id,
                        "image_path": image_path,
                        "image_index": image_counter,
                        "format": ext
                    })
            
        except Exception as e:
            print(f"Error extracting Word images: {str(e)}")
        
        return extracted_images
    
    def _extract_from_xlsx(self, file_path: str, output_folder: str,
                          questions_with_images: List[Dict]) -> List[Dict]:
        """Extract images from Excel file"""
        extracted_images = []
        
        try:
            from openpyxl.drawing.image import Image as OpenpyxlImage
            
            workbook = openpyxl.load_workbook(file_path)
            image_counter = 0
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                
                # Check for embedded images
                for image in sheet._images:
                    image_counter += 1
                    
                    # Get image data
                    image_data = image.ref.getvalue()
                    
                    # Assign to question
                    question_id = self._assign_image_to_question(
                        0, image_counter, questions_with_images
                    )
                    
                    # Save image
                    image_filename = f"{question_id}.png"
                    image_path = os.path.join(output_folder, image_filename)
                    
                    with open(image_path, "wb") as img_file:
                        img_file.write(image_data)
                    
                    extracted_images.append({
                        "question_id": question_id,
                        "image_path": image_path,
                        "sheet": sheet_name,
                        "image_index": image_counter,
                        "format": "png"
                    })
            
        except Exception as e:
            print(f"Error extracting Excel images: {str(e)}")
        
        return extracted_images
    
    def _assign_image_to_question(self, page_num: int, img_index: int, 
                                 questions_with_images: List[Dict]) -> str:
        """
        Assign extracted image to the most likely question
        This is a simplified implementation - in practice, you might use
        more sophisticated methods to match images to questions
        """
        if questions_with_images:
            # Simple assignment: distribute images evenly among questions
            question_index = (page_num + img_index) % len(questions_with_images)
            return questions_with_images[question_index].get("question_id", f"Q{question_index + 1:04d}")
        else:
            # Fallback: create generic question ID
            return f"Q{self.image_counter + 1:04d}"
    
    def process_extracted_images(self, image_list: List[Dict]) -> Dict[str, Any]:
        """
        Process and organize extracted images
        """
        processed_images = {}
        
        for img_info in image_list:
            question_id = img_info["question_id"]
            
            if question_id not in processed_images:
                processed_images[question_id] = []
            
            processed_images[question_id].append({
                "path": img_info["image_path"],
                "format": img_info["format"],
                "metadata": {
                    "page_number": img_info.get("page_number"),
                    "image_index": img_info.get("image_index"),
                    "sheet": img_info.get("sheet")
                }
            })
        
        return {
            "total_images": len(image_list),
            "questions_with_images": len(processed_images),
            "images_by_question": processed_images
        }
    
    def create_image_manifest(self, output_folder: str, processed_images: Dict) -> str:
        """
        Create a manifest file listing all extracted images
        """
        manifest_path = os.path.join(output_folder, "image_manifest.json")
        
        import json
        from datetime import datetime
        with open(manifest_path, 'w', encoding='utf-8') as f:
            json.dump({
                "generated_at": datetime.now().isoformat(),
                "total_images": processed_images["total_images"],
                "questions_with_images": processed_images["questions_with_images"],
                "image_details": processed_images["images_by_question"]
            }, f, indent=2, ensure_ascii=False)
        
        return manifest_path